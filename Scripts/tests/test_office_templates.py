"""Verify generated Office payloads, without adding product runtime dependencies.

Structural and regeneration checks use the standard library. Optional document
libraries additionally prove each blank document can be edited and saved.
"""

import base64
import importlib.util
from io import BytesIO
from pathlib import Path
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from urllib.parse import urlsplit
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile


ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "Scripts/generate-office-templates.py"
SWIFT = ROOT / "Shared/OfficeDocumentData.swift"
TEMPORARY = ROOT / ".build/Temporary/office-defaults-assets"
SPEC = importlib.util.spec_from_file_location("office_templates", SCRIPT)
GENERATOR = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(GENERATOR)
NS = {
    "w": GENERATOR.WORD, "s": GENERATOR.SHEET, "p": GENERATOR.PRESENTATION,
    "a": GENERATOR.DRAWING, "r": GENERATOR.REL,
    "rel": GENERATOR.PACKAGE_REL, "ct": GENERATOR.CONTENT_TYPES,
}


def available(module):
    return importlib.util.find_spec(module) is not None


def tearDownModule():
    if TEMPORARY.is_dir() and not any(TEMPORARY.iterdir()):
        TEMPORARY.rmdir()


class OfficeTemplateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.documents = GENERATOR.documents()

    def part(self, extension, name):
        with ZipFile(BytesIO(self.documents[extension])) as archive:
            return ET.fromstring(archive.read(name))

    def test_generated_swift_contains_the_reproducible_packages(self):
        self.assertEqual(SWIFT.read_text(), GENERATOR.swift_source())
        self.assertEqual(self.documents, GENERATOR.documents())
        encoded = re.findall(
            r'static let (docx|xlsx|pptx): Data = Data\(base64Encoded: """\n(.*?)\n\s*"""',
            SWIFT.read_text(), re.DOTALL,
        )
        self.assertEqual(len(encoded), 3)
        for extension, content in encoded:
            self.assertEqual(base64.b64decode(content), self.documents[extension])

    def test_regeneration_cli_checks_and_detects_stale_output(self):
        TEMPORARY.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TEMPORARY) as directory:
            output = Path(directory) / "OfficeDocumentData.swift"
            command = [sys.executable, str(SCRIPT), "--output", str(output)]
            subprocess.run(command, check=True, capture_output=True)
            self.assertEqual(output.read_bytes(), SWIFT.read_bytes())
            subprocess.run(command + ["--check"], check=True, capture_output=True)
            output.write_text("stale\n")
            result = subprocess.run(command + ["--check"], capture_output=True, text=True)
            self.assertEqual(result.returncode, 1)
            self.assertIn("Office document data differs", result.stderr)
            self.assertEqual(output.read_text(), "stale\n")

    def test_zip_parts_have_fixed_metadata_and_valid_xml(self):
        for extension, content in self.documents.items():
            with self.subTest(extension=extension), ZipFile(BytesIO(content)) as archive:
                self.assertIsNone(archive.testzip())
                self.assertEqual(archive.namelist(), sorted(set(archive.namelist())))
                self.assertEqual(archive.comment, b"")
                for info in archive.infolist():
                    self.assertEqual(info.date_time, GENERATOR.ZIP_TIME)
                    self.assertEqual(info.compress_type, ZIP_DEFLATED)
                    self.assertEqual(info.create_system, 3)
                    self.assertEqual(info.external_attr, 0o100644 << 16)
                    self.assertEqual(info.flag_bits & 1, 0)
                    self.assertEqual(info.extra, b"")
                    self.assertEqual(info.comment, b"")
                    self.assertEqual(posixpath.normpath(info.filename), info.filename)
                    ET.fromstring(archive.read(info))

    def test_all_relationship_targets_and_references_resolve_inside_the_package(self):
        for extension, content in self.documents.items():
            with self.subTest(extension=extension), ZipFile(BytesIO(content)) as archive:
                names = set(archive.namelist())
                relations = {}
                for name in names:
                    if not name.endswith(".rels"):
                        continue
                    source = "" if name == "_rels/.rels" else name.replace("_rels/", "")[:-5]
                    if source:
                        self.assertIn(source, names)
                    targets = {}
                    for relation in ET.fromstring(archive.read(name)):
                        identifier = relation.get("Id")
                        self.assertNotIn(identifier, targets)
                        self.assertIsNone(relation.get("TargetMode"))
                        target = relation.get("Target")
                        self.assertFalse(urlsplit(target).scheme)
                        self.assertFalse(urlsplit(target).netloc)
                        resolved = posixpath.normpath(posixpath.join(posixpath.dirname(source), target))
                        self.assertIn(resolved, names)
                        targets[identifier] = resolved
                    relations[source] = targets
                self.assertEqual(len(relations[""]), 1)
                for name in names - {"[Content_Types].xml"}:
                    if name.endswith(".rels"):
                        continue
                    for element in ET.fromstring(archive.read(name)).iter():
                        for attribute, identifier in element.attrib.items():
                            if attribute.startswith("{" + GENERATOR.REL + "}"):
                                self.assertIn(identifier, relations[name])

    def test_content_types_cover_every_part_without_macros_or_personal_properties(self):
        main_types = {
            "docx": ("word/document.xml", "wordprocessingml.document.main+xml"),
            "xlsx": ("xl/workbook.xml", "spreadsheetml.sheet.main+xml"),
            "pptx": ("ppt/presentation.xml", "presentationml.presentation.main+xml"),
        }
        for extension, content in self.documents.items():
            with self.subTest(extension=extension), ZipFile(BytesIO(content)) as archive:
                names = set(archive.namelist())
                manifest = ET.fromstring(archive.read("[Content_Types].xml"))
                overrides = {
                    entry.get("PartName")[1:]: entry.get("ContentType")
                    for entry in manifest.findall("ct:Override", NS)
                }
                self.assertEqual(set(overrides), {
                    name for name in names if name != "[Content_Types].xml" and not name.endswith(".rels")
                })
                main, kind = main_types[extension]
                self.assertEqual(overrides[main], GENERATOR.MIME + kind)
                root_relationships = ET.fromstring(archive.read("_rels/.rels"))
                self.assertEqual(len(root_relationships), 1)
                self.assertEqual(root_relationships[0].get("Target"), main)
                self.assertEqual(root_relationships[0].get("Type"), GENERATOR.REL + "/officeDocument")
                for name in names:
                    self.assertFalse(any(value in name.lower() for value in (
                        "vbaproject", "docprops", "thumbnail", "customxml", "embeddings", "/media/",
                    )))
                    payload = archive.read(name).lower()
                    for forbidden in (b"macroenabled", b"/users/", b"file:", b"lastmodifiedby", b"dc:creator"):
                        self.assertNotIn(forbidden, payload)

    def test_word_is_one_blank_a4_section(self):
        document = self.part("docx", "word/document.xml")
        self.assertEqual(len(document.findall("w:body/w:p", NS)), 1)
        self.assertEqual(document.findall(".//w:t", NS), [])
        self.assertEqual(document.findall(".//w:tbl", NS), [])
        self.assertEqual(document.findall(".//w:br", NS), [])
        size = document.find("w:body/w:sectPr/w:pgSz", NS)
        self.assertEqual(size.get("{" + GENERATOR.WORD + "}w"), "11906")
        self.assertEqual(size.get("{" + GENERATOR.WORD + "}h"), "16838")

    def test_spreadsheet_is_one_empty_sheet(self):
        workbook = self.part("xlsx", "xl/workbook.xml")
        sheets = workbook.findall("s:sheets/s:sheet", NS)
        self.assertEqual(len(sheets), 1)
        self.assertEqual(sheets[0].get("name"), "Sheet1")
        sheet = self.part("xlsx", "xl/worksheets/sheet1.xml")
        self.assertEqual(len(sheet.find("s:sheetData", NS)), 0)
        self.assertEqual(sheet.findall(".//s:c", NS), [])

    def test_presentation_is_one_empty_widescreen_slide_with_editable_layout(self):
        presentation = self.part("pptx", "ppt/presentation.xml")
        self.assertEqual(len(presentation.findall("p:sldIdLst/p:sldId", NS)), 1)
        size = presentation.find("p:sldSz", NS)
        self.assertEqual(int(size.get("cx")) * 9, int(size.get("cy")) * 16)
        for name in ("ppt/slides/slide1.xml", "ppt/slideLayouts/slideLayout1.xml", "ppt/slideMasters/slideMaster1.xml"):
            part = self.part("pptx", name)
            self.assertEqual(part.findall(".//p:sp", NS), [])
            self.assertEqual(part.findall(".//a:t", NS), [])
            self.assertIsNotNone(part.find("p:cSld/p:spTree", NS))
        layout = self.part("pptx", "ppt/slideLayouts/slideLayout1.xml")
        self.assertEqual(layout.get("type"), "blank")
        theme = self.part("pptx", "ppt/theme/theme1.xml")
        for name in ("fillStyleLst", "lnStyleLst", "effectStyleLst", "bgFillStyleLst"):
            self.assertEqual(len(theme.find("a:themeElements/a:fmtScheme/a:" + name, NS)), 3)

    @unittest.skipUnless(available("docx"), "python-docx is an optional validation dependency")
    def test_word_opens_and_can_be_edited_and_saved(self):
        from docx import Document

        document = Document(BytesIO(self.documents["docx"]))
        self.assertEqual([paragraph.text for paragraph in document.paragraphs], [""])
        self.assertEqual(len(document.sections), 1)
        document.paragraphs[0].text = "Editable"
        output = BytesIO()
        document.save(output)
        output.seek(0)
        self.assertEqual(Document(output).paragraphs[0].text, "Editable")

    @unittest.skipUnless(available("openpyxl"), "openpyxl is an optional validation dependency")
    def test_spreadsheet_opens_and_can_add_data_and_sheets(self):
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(self.documents["xlsx"]))
        self.assertEqual(workbook.sheetnames, ["Sheet1"])
        self.assertEqual(list(workbook.active.values), [])
        workbook.active["A1"] = "Editable"
        workbook.create_sheet("Sheet2")
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        reopened = load_workbook(output)
        self.assertEqual(reopened.active["A1"].value, "Editable")
        self.assertEqual(reopened.sheetnames, ["Sheet1", "Sheet2"])

    @unittest.skipUnless(available("pptx"), "python-pptx is an optional validation dependency")
    def test_presentation_opens_and_can_add_text_and_slides(self):
        from pptx import Presentation

        presentation = Presentation(BytesIO(self.documents["pptx"]))
        self.assertEqual(len(presentation.slides), 1)
        self.assertEqual(len(presentation.slides[0].shapes), 0)
        self.assertEqual(len(presentation.slide_masters), 1)
        self.assertEqual(len(presentation.slide_layouts), 1)
        layout = presentation.slide_layouts[0]
        self.assertEqual(layout.name, "Blank")
        presentation.slides[0].shapes.add_textbox(0, 0, 1000000, 1000000).text = "Editable"
        presentation.slides.add_slide(layout)
        output = BytesIO()
        presentation.save(output)
        output.seek(0)
        reopened = Presentation(output)
        self.assertEqual(len(reopened.slides), 2)
        self.assertEqual(reopened.slides[0].shapes[0].text, "Editable")
        self.assertEqual(len(reopened.slides[1].shapes), 0)

    @unittest.skipUnless(shutil.which("swiftc"), "Swift compiler is unavailable")
    def test_swift_data_decodes_to_the_exact_generated_documents(self):
        TEMPORARY.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TEMPORARY) as directory:
            fixture = Path(directory)
            main = fixture / "main.swift"
            main.write_text('import Foundation\n'
                            'for data in [OfficeDocumentData.docx, OfficeDocumentData.xlsx, OfficeDocumentData.pptx] {\n'
                            '    print(data.base64EncodedString())\n}\n')
            binary = fixture / "office-data"
            subprocess.run(
                ["swiftc", "-module-cache-path", str(fixture / "module-cache"), str(SWIFT), str(main), "-o", str(binary)],
                check=True, capture_output=True, text=True,
            )
            output = subprocess.run([str(binary)], check=True, capture_output=True, text=True)
            self.assertEqual(
                [base64.b64decode(line) for line in output.stdout.splitlines()],
                list(self.documents.values()),
            )


if __name__ == "__main__":
    unittest.main()
